from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
import os

OUTPUT_PATH = "docs/sablona_zaverecne_zadanie.docx"

def set_doc_style(doc):
    section = doc.sections[0]
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)

def add_space(doc, n=1):
    for _ in range(n):
        doc.add_paragraph("")

def main():
    os.makedirs("docs", exist_ok=True)

    doc = Document()
    set_doc_style(doc)

    title = doc.add_heading("ZÁVEREČNÉ ZADANIE", level=0)
    title.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

    doc.add_paragraph("Predmet: ")
    doc.add_paragraph("Meno a priezvisko: ")
    doc.add_paragraph("Študentský email: ")
    doc.add_paragraph("Dátum odovzdania: ")

    add_space(doc, 1)

    doc.add_heading("1. Úvod", level=1)
    add_space(doc, 3)

    doc.add_heading("2. Dáta", level=1)
    add_space(doc, 3)

    doc.add_heading("3. Metodika", level=1)
    doc.add_heading("3.1 Sliding window", level=2)
    add_space(doc, 2)
    doc.add_heading("3.2 Modely", level=2)
    add_space(doc, 2)
    doc.add_heading("3.3 Metriky (MSE, RMSE, MAPE)", level=2)
    add_space(doc, 2)

    doc.add_heading("4. Výsledky", level=1)
    add_space(doc, 2)

    table1 = doc.add_table(rows=1 + 4, cols=4)
    table1.style = "Table Grid"
    headers = ["Model", "MSE", "RMSE", "MAPE"]
    for i, h in enumerate(headers):
        table1.rows[0].cells[i].text = h

    add_space(doc, 2)

    doc.add_heading("5. Feature engineering (pred vs. po)", level=1)
    add_space(doc, 2)

    table2 = doc.add_table(rows=1 + 4, cols=4)
    table2.style = "Table Grid"
    for i, h in enumerate(headers):
        table2.rows[0].cells[i].text = h

    add_space(doc, 2)

    doc.add_heading("6. Záver", level=1)
    add_space(doc, 3)

    doc.add_heading("Prílohy (voliteľné)", level=1)
    add_space(doc, 2)

    doc.save(OUTPUT_PATH)
    print(f"Hotovo: {OUTPUT_PATH}")

if __name__ == "__main__":
    main()